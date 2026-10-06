"""Natijani Excel faylga chiqarish: devonxona yoki ijro nazorati bo'limi o'z jadvaliga ko'chira oladi."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

TYPES = {"sana": "aniq sana", "nisbiy": "hujjat sanasidan", "shartli": "shartli", "davriy": "davriy", "doimiy": "doimiy",
         "boshlanish": "boshlanish sanasi", "matn": "matn", "yo'q": "ko'rsatilmagan"}


def _fmt(d: str | None) -> str:
    return ".".join(reversed(d.split("-"))) if d else ""


def to_xlsx(result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Topshiriqlar"
    head = ["№", "Topshiriq", "Ijrochi", "Muddat", "Muddat turi", "Sana", "Natija shakli", "Manba"]
    ws.append(head)
    for t in result.get("tasks", []):
        d = t.get("deadline") or {}
        ws.append([t.get("n"), t.get("text"), "\n".join(t.get("executors", [])), d.get("text", ""),
                   TYPES.get(d.get("type"), d.get("type", "")), _fmt(d.get("date")), t.get("result", ""), t.get("source", "")])
    widths = [5, 60, 32, 26, 16, 12, 18, 18]
    for col, w in zip("ABCDEFGH", widths):
        ws.column_dimensions[col].width = w
    _style(ws)

    rq = result.get("requisites", {})
    ws2 = wb.create_sheet("Rekvizitlar")
    ws2.append(["Maydon", "Qiymat"])
    rows = [("Hujjat turi", rq.get("doc_type")), ("Jo'natuvchi", rq.get("sender")), ("Hujjat raqami", rq.get("number")),
            ("Sana", _fmt(rq.get("date"))), ("Imzolagan shaxs", rq.get("signatory")),
            ("Lavozimi", rq.get("signatory_position")), ("Kimga", rq.get("addressee")), ("Mavzu", rq.get("title")),
            ("Hujjat ijrochisi", rq.get("preparer")), ("Ijro nazorati", "; ".join(result.get("control", [])))]
    for k, v in rows:
        ws2.append([k, v or ""])
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 90
    _style(ws2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _style(ws):
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF6")
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
